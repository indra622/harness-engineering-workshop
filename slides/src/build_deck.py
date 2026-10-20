"""Build the workshop deck as an editable PowerPoint file.

Usage (from the project root):
    python3 slides/src/build_deck.py --set representative \
        --out slides/previews/representative-v01.pptx

--set representative builds the six proof slides (S1, S5, S11, S17, S26, S34).
--set all builds every slide ID that has a definition in deck/slides.py, in
v10 order. Slide IDs are never renumbered: a sample keeps its source IDs.
"""

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE))

from deck import kit, notes  # noqa: E402
from deck import theme as T  # noqa: E402
from deck.slides import REPRESENTATIVE, SLIDES  # noqa: E402


def build(ids, out):
    prs = Presentation()
    prs.slide_width = Inches(T.SLIDE_W)
    prs.slide_height = Inches(T.SLIDE_H)
    blank = prs.slide_layouts[6]
    for sid in ids:
        slide = prs.slides.add_slide(blank)
        page = kit.Page(slide, sid)
        h = notes.heading(ROOT, sid)
        SLIDES[sid](page, h)
        slide.notes_slide.notes_text_frame.text = notes.notes_for(ROOT, sid)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.core_properties.title = "하네스 엔지니어링 워크샵"
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--set", choices=["representative", "all"], default="representative")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    ids = REPRESENTATIVE if a.set == "representative" else sorted(SLIDES)
    out = build(ids, Path(a.out))
    print(f"built {len(ids)} slides: {', '.join(f'S{i}' for i in ids)} -> {a.out}")


if __name__ == "__main__":
    main()
