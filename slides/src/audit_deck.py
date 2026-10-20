"""Audit a rendered deck and write an evidence JSON.

Usage (from the project root, after build and render):
    python3 slides/src/audit_deck.py slides/previews/representative-v01.pptx

Reads the PPTX, the PDF and the PNG pages produced by render.sh and checks
what a reviewer would otherwise check by eye: source slide IDs and titles,
font sizes and family, text fit with real font metrics, safe-area bounds,
unintended overlaps, inherited theme effects, speaker notes, embedded PDF
fonts, PDF text, page count and file hashes. Exit code 1 if any check fails.
"""

import hashlib
import json
import re
import subprocess
import sys
from itertools import combinations
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE))

from deck import kit, notes  # noqa: E402
from deck import theme as T  # noqa: E402

LEAK = re.compile(r"/Users/|/Volumes/|/System/|/private/|session_[0-9A-Za-z]{6,}")


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def rel(p):
    return str(Path(p).resolve().relative_to(ROOT))


def inches(v):
    return Emu(v).inches


def para_runs(p):
    runs = []
    for r in p.runs:
        size = r.font.size.pt if r.font.size else None
        runs.append((r.text, size, bool(r.font.bold), r.font.name))
    return runs


def audit(pptx_path):
    pptx_path = Path(pptx_path)
    name = pptx_path.stem
    pdf = pptx_path.with_suffix(".pdf")
    png_dir = pptx_path.parent / name
    prs = Presentation(pptx_path)
    problems = []
    slides_out = []
    font_sizes = []

    for idx, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        texts = [s for s in shapes if s.has_text_frame and s.text_frame.text.strip()]
        meta = next(s for s in shapes if s.name.startswith("T:") and "meta-left" in s.name)
        sid = int(re.match(r"S(\d+)", meta.text_frame.text.strip()).group(1))
        head = notes.heading(ROOT, sid)
        title = next(s for s in shapes if s.name.endswith(":title"))
        title_text = " ".join(p.text for p in title.text_frame.paragraphs)
        if title_text != head["title"]:
            problems.append(f"S{sid}: title '{title_text}' != v10 '{head['title']}'")

        fit_fail, small, wrong_font, out_of_bounds, styled = [], [], [], [], []
        for s in shapes:
            x, y = inches(s.left), inches(s.top)
            w, h = inches(s.width), inches(s.height)
            if (x < T.SAFE - 1e-3 or y < T.SAFE - 1e-3 or x + w > T.SLIDE_W - T.SAFE + 1e-3
                    or y + h > T.SLIDE_H - T.SAFE + 1e-3):
                out_of_bounds.append(s.name)
            if s._element.find(qn("p:style")) is not None:
                styled.append(s.name)
            if not (s.has_text_frame and s.text_frame.text.strip()):
                continue
            tf = s.text_frame
            ml, mr = inches(tf.margin_left), inches(tf.margin_right)
            mt, mb = inches(tf.margin_top), inches(tf.margin_bottom)
            total_h = 0.0
            for p in tf.paragraphs:
                runs = para_runs(p)
                if not runs:
                    continue
                width = 0.0
                line_pt = 0
                for text, size, bold, font in runs:
                    font_sizes.append(size)
                    is_diagram = s.name.startswith("N:")
                    floor = T.SIZE["label"] if is_diagram else T.SIZE["footnote"]
                    if size is None or size < floor:
                        small.append(f"{s.name}:{text[:12]}:{size}")
                    if font != T.FONT:
                        wrong_font.append(s.name)
                    width += kit.text_width_in(text, size, bold)
                    line_pt = max(line_pt, size)
                if width > w - ml - mr + 0.01:
                    fit_fail.append(f"{s.name}: line {width:.2f}in > box {w - ml - mr:.2f}in")
                total_h += line_pt * kit.LINE_HEIGHT / 72
            if total_h > h - mt - mb + 0.02:
                fit_fail.append(f"{s.name}: text {total_h:.2f}in > box {h - mt - mb:.2f}in")

        # unintended overlaps between nodes and free text (containers, bands, connectors excluded)
        boxes = [(s.name, inches(s.left), inches(s.top), inches(s.width), inches(s.height))
                 for s in shapes if s.name[:2] in ("N:", "T:")]
        overlaps = []
        for a, b in combinations(boxes, 2):
            ox = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            oy = min(a[2] + a[4], b[2] + b[4]) - max(a[2], b[2])
            if ox > 0.02 and oy > 0.02:
                overlaps.append(f"{a[0]} x {b[0]}")

        nt = slide.notes_slide.notes_text_frame.text
        plan = (ROOT / notes.PLAN).read_text(encoding="utf-8")
        v10_block = notes._plain(notes._block(plan, sid))
        note_checks = {
            "has_source_header": f"원본 v10 S{sid}." in nt,
            "has_full_v10_block": v10_block in nt,
            "has_brief_entry": "[제작 브리프]" in nt and "제작 검수 포인트" in nt,
            "has_time": f"{head['start']}분부터 {head['end']}분까지" in nt,
        }
        all_text = nt + "\n".join(s.text_frame.text for s in texts)
        leaks = sorted(set(LEAK.findall(all_text)))

        for label, items in [("text overflow", fit_fail), ("font below floor", small),
                             ("font not Nanum Gothic", wrong_font),
                             ("outside 0.4in safe area", out_of_bounds),
                             ("theme style left on shape", styled), ("overlap", overlaps),
                             ("host path in text or notes", leaks)]:
            if items:
                problems.append(f"S{sid}: {label}: {items[:5]}")
        for k, ok in note_checks.items():
            if not ok:
                problems.append(f"S{sid}: notes missing {k}")

        slides_out.append({
            "page": idx, "source_slide_id": f"S{sid}", "title": title_text,
            "minutes": head["minutes"], "shapes": len(shapes),
            "text_shapes": len(texts),
            "connectors": sum(1 for s in shapes if s.name.startswith("L:")),
            "diagram_nodes": sum(1 for s in shapes if s.name.startswith("N:")),
            "characters_on_slide": sum(len(s.text_frame.text) for s in texts),
            "notes_characters": len(nt), "notes_checks": note_checks,
        })

    # PDF and render checks
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    fonts = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    font_names = sorted({ln.split()[0].split("+")[-1] for ln in fonts.splitlines()[2:] if ln.strip()})
    pdf_text = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True).stdout
    pngs = sorted(png_dir.glob("page-*.png"))
    if pages != len(prs.slides):
        problems.append(f"PDF pages {pages} != slides {len(prs.slides)}")
    if any("Nanum" not in f for f in font_names):
        problems.append(f"PDF embeds non-Nanum fonts: {font_names}")
    for s in slides_out:
        if s["title"] not in pdf_text.replace("\n", " "):
            problems.append(f"{s['source_slide_id']}: title not found in PDF text")
    if len(pngs) != len(prs.slides):
        problems.append(f"PNG pages {len(pngs)} != slides {len(prs.slides)}")
    meta_leak = LEAK.findall(" ".join(filter(None, [prs.core_properties.title,
                                                   prs.core_properties.author,
                                                   prs.core_properties.last_modified_by])))
    if meta_leak:
        problems.append(f"host path in document properties: {meta_leak}")

    evidence = {
        "deck": rel(pptx_path),
        "phase": "representative proof, 6 of 35 source slides; not the finished deck",
        "slide_size_in": [T.SLIDE_W, T.SLIDE_H],
        "slides": slides_out,
        "font_family": T.FONT,
        "font_size_min_pt": min(f for f in font_sizes if f),
        "diagram_label_floor_pt": T.SIZE["label"],
        "pdf": {"file": rel(pdf), "pages": pages, "embedded_fonts": font_names},
        "render": {"pages": len(pngs), "files": [rel(p) for p in pngs]},
        "hashes_sha256": {rel(p): sha(p) for p in [pptx_path, pdf, *pngs]},
        "checks": {
            "source_ids_and_titles_match_v10": not any("title" in p for p in problems),
            "text_fits_boxes": not any("overflow" in p for p in problems),
            "fonts_at_or_above_floor": not any("floor" in p for p in problems),
            "inside_safe_area": not any("safe area" in p for p in problems),
            "no_unintended_overlap": not any("overlap" in p for p in problems),
            "no_inherited_theme_effects": not any("theme style" in p for p in problems),
            "notes_populated": not any("notes missing" in p for p in problems),
            "pdf_fonts_nanum_only": not any("non-Nanum" in p for p in problems),
            "no_host_paths": not any("host path" in p for p in problems),
        },
        "problems": problems,
        "visual_review": "page PNGs viewed by the worker; see report for findings",
    }
    out = pptx_path.parent / f"{name}-evidence.json"
    out.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return out, problems


if __name__ == "__main__":
    out, problems = audit(sys.argv[1])
    print(f"evidence: {rel(out)}")
    for p in problems:
        print("FAIL", p)
    print("ok" if not problems else f"{len(problems)} problem(s)")
    sys.exit(1 if problems else 0)
