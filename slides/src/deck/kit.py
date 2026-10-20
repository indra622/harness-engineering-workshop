"""Drawing kit: native, editable PowerPoint shapes for the workshop deck.

Every helper places one editable shape (text box, rounded rectangle,
ellipse or connector). Shape names carry a role prefix that the audit uses:

  T:  free text            N:  node (labelled box)
  C:  container (may hold N/T shapes inside it)
  B:  band (question strip, lane band)   L:  connector

Text is laid out with explicit line breaks. `fits()` measures each line
with the real font metrics so that labels never depend on automatic wrap.
"""

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import theme as T

LINE_HEIGHT = 1.2      # line height multiple used for fit checks
INSET = 0.06           # text inset in inches (left/right)
INSET_V = 0.03


from .metrics import text_width_in  # noqa: E402,F401  (re-exported for the PPTX code)


def rgb(hex6):
    return RGBColor.from_string(hex6)


# ---------------------------------------------------------------- text runs

def _set_font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = T.FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", T.FONT)


def fill_text(tf, paragraphs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
              line_spacing=1.08):
    """paragraphs: list of lists of (text, size, bold, color) runs."""
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(INSET)
    tf.margin_top = tf.margin_bottom = Inches(INSET_V)
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            _set_font(r, size, bold, color)


def runs_of(text, size, bold=False, color=T.COLORS["ink"]):
    """Split explicit '\n' lines into one paragraph each."""
    return [[(line, size, bold, color)] for line in text.split("\n")]


# ---------------------------------------------------------------- registry

class Page:
    """Wraps one slide and records every placed element for the layout report."""

    def __init__(self, slide, sid):
        self.slide = slide
        self.sid = sid
        self.items = []
        self._n = 0

    def _name(self, prefix, hint):
        self._n += 1
        return f"{prefix}:{self.sid}:{self._n:02d}:{hint}"[:120]

    def record(self, shape, role, paragraphs=None, parent=None):
        self.items.append({
            "name": shape.name, "role": role, "parent": parent,
            "x": Emu(shape.left).inches, "y": Emu(shape.top).inches,
            "w": Emu(shape.width).inches, "h": Emu(shape.height).inches,
            "paragraphs": paragraphs or [],
        })
        return shape


def _drop_theme_style(shp):
    """Remove the theme style reference so no inherited shadow or color applies.
    Fill, line and text color are always set explicitly instead."""
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)


def _shape_style(shp, fill, line, width, dash=False):
    _drop_theme_style(shp)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(width)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    shp.shadow.inherit = False


def _round(shp, w, h, radius=0.09):
    try:
        shp.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    except IndexError:
        pass


def text(page, x, y, w, h, paragraphs, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE, role="T", parent=None, hint="text"):
    tb = page.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = page._name(role, hint)
    fill_text(tb.text_frame, paragraphs, align=align, anchor=anchor)
    return page.record(tb, role, paragraphs, parent)


STYLES = {
    # fill, line color, line width, dashed, text color, bold
    "base":     ("surface", "ink", "default", False, "ink", False),
    "focus":    ("focus_fill", "primary", "focus", False, "ink", False),
    "gate":     ("surface", "warning", "gate", False, "ink", True),
    "retained": ("canvas", "muted", "focus", True, "ink", False),
    "out":      (None, "muted", "default", True, "muted", False),
    "chip":     ("band", None, "default", False, "ink", False),
    "note":     (None, "muted", "hair", True, "muted", False),
    "caution":  ("surface", "warning", "focus", True, "warning", True),
    "tile":     ("surface", "hairline", "hair", False, "ink", False),
    "scope":    (None, "primary", "focus", False, "ink", False),
    "plain":    ("surface", None, "default", False, "ink", False),
}


def node(page, x, y, w, h, label, style="base", size=T.SIZE["label"],
         parent=None, hint=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         align=PP_ALIGN.CENTER, paragraphs=None):
    fill, line, lw, dash, tcolor, bold = STYLES[style]
    shp = page.slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = page._name("N", hint or label.split("\n")[0])
    _shape_style(shp, T.COLORS[fill] if fill else None,
                 T.COLORS[line] if line else None, T.LINE[lw], dash)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        _round(shp, w, h)
    paras = paragraphs or runs_of(label, size, bold, T.COLORS[tcolor])
    fill_text(shp.text_frame, paras, align=align)
    return page.record(shp, "N", paras, parent)


def container(page, x, y, w, h, style="focus", hint="container", fill=None):
    fill_key, line, lw, dash, _, _ = STYLES[style]
    shp = page.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = page._name("C", hint)
    use_fill = fill if fill is not None else fill_key
    _shape_style(shp, T.COLORS[use_fill] if use_fill else None,
                 T.COLORS[line] if line else None, T.LINE[lw], dash)
    _round(shp, w, h, 0.12)
    return page.record(shp, "C")


def band(page, x, y, w, h, color="band", paragraphs=None, hint="band",
         line=None, line_w="hair", align=PP_ALIGN.LEFT):
    shp = page.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                      Inches(w), Inches(h))
    shp.name = page._name("B", hint)
    _shape_style(shp, T.COLORS[color], T.COLORS[line] if line else None, T.LINE[line_w])
    if paragraphs:
        fill_text(shp.text_frame, paragraphs, align=align)
        shp.text_frame.margin_left = Inches(0.3)
    return page.record(shp, "B", paragraphs or [])


def _arrowhead(connector, color, width, dash=False):
    _drop_theme_style(connector)
    ln = connector.line
    ln.color.rgb = rgb(color)
    ln.width = Pt(width)
    if dash:
        ln.dash_style = MSO_LINE.DASH
    el = ln._get_or_add_ln()
    tail = el.find(qn("a:tailEnd"))
    if tail is None:
        tail = etree.SubElement(el, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


# connection point indices for rectangles: 0 top, 1 left, 2 bottom, 3 right
SIDE = {"top": 0, "left": 1, "bottom": 2, "right": 3}


def arrow(page, a, b, sa="right", sb="left", color="ink", width="connector",
          dash=False, elbow=False):
    """Connector glued to shape a (side sa) and shape b (side sb), arrowhead at b."""
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    c = page.slide.shapes.add_connector(kind, 0, 0, 0, 0)
    c.begin_connect(a, SIDE[sa])
    c.end_connect(b, SIDE[sb])
    c.name = page._name("L", f"{a.name.split(':')[-1]}->{b.name.split(':')[-1]}")
    _arrowhead(c, T.COLORS[color], T.LINE[width], dash)
    return page.record(c, "L")


def free_arrow(page, x1, y1, x2, y2, color="ink", width="connector", dash=False):
    c = page.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                        Inches(x2), Inches(y2))
    c.name = page._name("L", "free")
    _arrowhead(c, T.COLORS[color], T.LINE[width], dash)
    return page.record(c, "L")


def ellipse(page, x, y, d, label, style="base", parent=None):
    return node(page, x, y, d, d, label, style=style, parent=parent,
                shape=MSO_SHAPE.OVAL, hint=label)
